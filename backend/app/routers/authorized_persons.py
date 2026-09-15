from fastapi import APIRouter, Depends, HTTPException, UploadFile, File
from sqlalchemy.orm import Session
from typing import List
import csv
from io import StringIO, BytesIO
from fastapi.responses import StreamingResponse
from openpyxl import Workbook, load_workbook
from .. import crud, schemas
from ..database import SessionLocal
from .admin import verify_admin

router = APIRouter()

def get_db():
    db = SessionLocal()
    try:
        yield db
    finally:
        db.close()

# نگاشت عنوان دقیق ستون‌های فایل اکسل به نام فیلد در دیتابیس
COLUMN_MAP = {
    "کدملی": "national_code",
    "کدپرسنلی": "personnel_code",
    "نام": "first_name",
    "نام خانوادگی": "last_name",
    "محل خدمت": "service_location",
    "وضعیت خدمت": "service_status",
}
REQUIRED_COLUMNS = ["کدملی", "نام", "نام خانوادگی"]

# ========== لیست و مدیریت تکی ==========
@router.get("/lookup/{national_code}", response_model=schemas.AuthorizedPerson)
def lookup_authorized_person(national_code: str, db: Session = Depends(get_db)):
    """
    روت عمومی (بدون نیاز به توکن ادمین): برای جست‌وجوی مشخصات یک نفر بر اساس کدملی
    در فرم عمومی ثبت‌درخواست استفاده می‌شود.
    """
    person = crud.get_authorized_person_by_national_code(db, national_code)
    if not person:
        raise HTTPException(
            status_code=404,
            detail="کد ملی شما در سامانه ثبت نشده است، برای حل مشکل به امور مالی اداره مراجعه نمایید"
        )
    return person

@router.get("", response_model=List[schemas.AuthorizedPerson])
def list_authorized_persons(
    skip: int = 0,
    limit: int = 2000,
    db: Session = Depends(get_db),
    admin: str = Depends(verify_admin)
):
    return crud.get_authorized_persons(db, skip=skip, limit=limit)

@router.post("", response_model=schemas.AuthorizedPerson)
def create_authorized_person(
    person: schemas.AuthorizedPersonCreate,
    db: Session = Depends(get_db),
    admin: str = Depends(verify_admin)
):
    existing = crud.get_authorized_person_by_national_code(db, person.national_code)
    if existing:
        raise HTTPException(status_code=400, detail="این کدملی قبلاً در لیست ثبت شده است")
    return crud.create_authorized_person(db, person)

# ========== قالب نمونه ==========
@router.get("/template")
def download_template(admin: str = Depends(verify_admin)):
    wb = Workbook()
    ws = wb.active
    ws.title = "افراد مجاز"
    ws.append(["کدملی", "کدپرسنلی", "نام", "نام خانوادگی", "محل خدمت", "وضعیت خدمت"])
    ws.append(["0012345678", "12345678", "علی", "رضایی", "دفتر مرکزی", "شاغل"])
    output = BytesIO()
    wb.save(output)
    output.seek(0)
    return StreamingResponse(
        output,
        media_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
        headers={"Content-Disposition": "attachment; filename=authorized_persons_template.xlsx"}
    )

# ========== آپلود فایل اکسل ==========
@router.post("/upload", response_model=schemas.AuthorizedPersonUploadResult)
async def upload_authorized_persons(
    file: UploadFile = File(...),
    db: Session = Depends(get_db),
    admin: str = Depends(verify_admin)
):
    if not file.filename or not file.filename.lower().endswith((".xlsx", ".xls")):
        raise HTTPException(status_code=400, detail="فقط فایل اکسل (.xlsx) پذیرفته می‌شود")

    content = await file.read()
    try:
        wb = load_workbook(BytesIO(content), data_only=True)
    except Exception:
        raise HTTPException(status_code=400, detail="فایل اکسل قابل خواندن نیست. لطفاً از فرمت صحیح استفاده کنید")

    ws = wb.active
    rows_iter = ws.iter_rows(values_only=True)
    try:
        header_row = next(rows_iter)
    except StopIteration:
        raise HTTPException(status_code=400, detail="فایل خالی است")

    header_index = {}
    for idx, cell in enumerate(header_row):
        if cell is None:
            continue
        title = str(cell).strip()
        if title in COLUMN_MAP:
            header_index[COLUMN_MAP[title]] = idx

    missing = [c for c in REQUIRED_COLUMNS if COLUMN_MAP[c] not in header_index]
    if missing:
        raise HTTPException(
            status_code=400,
            detail=f"ستون‌های زیر در فایل یافت نشد: {', '.join(missing)}"
        )

    def get_cell(row, field):
        idx = header_index.get(field)
        if idx is None or idx >= len(row):
            return None
        val = row[idx]
        if val is None:
            return None
        val_str = str(val).strip()
        return val_str if val_str else None

    valid_rows = []
    errors = []
    total = 0
    for row_num, row in enumerate(rows_iter, start=2):
        if row is None or all(c is None for c in row):
            continue
        total += 1

        national_code = get_cell(row, "national_code")
        first_name = get_cell(row, "first_name")
        last_name = get_cell(row, "last_name")

        if not national_code:
            errors.append(f"ردیف {row_num}: کدملی خالی است - نادیده گرفته شد")
            continue
        if not first_name or not last_name:
            errors.append(f"ردیف {row_num}: نام یا نام خانوادگی خالی است - نادیده گرفته شد")
            continue

        valid_rows.append({
            "national_code": national_code,
            "personnel_code": get_cell(row, "personnel_code"),
            "first_name": first_name,
            "last_name": last_name,
            "service_location": get_cell(row, "service_location"),
            "service_status": get_cell(row, "service_status"),
        })

    created, updated = crud.bulk_upsert_authorized_persons(db, valid_rows)

    return schemas.AuthorizedPersonUploadResult(
        total_rows=total,
        created=created,
        updated=updated,
        errors=errors
    )

# ========== خروجی گرفتن ==========
@router.get("/export/csv")
def export_csv(db: Session = Depends(get_db), admin: str = Depends(verify_admin)):
    persons = crud.get_authorized_persons(db, limit=1000000)
    output = StringIO()
    writer = csv.writer(output)
    writer.writerow(["کدملی", "کدپرسنلی", "نام", "نام خانوادگی", "محل خدمت", "وضعیت خدمت"])
    for p in persons:
        writer.writerow([
            p.national_code,
            p.personnel_code or "",
            p.first_name,
            p.last_name,
            p.service_location or "",
            p.service_status or ""
        ])
    return StreamingResponse(
        iter([output.getvalue()]),
        media_type="text/csv",
        headers={"Content-Disposition": "attachment; filename=authorized_persons.csv"}
    )

@router.get("/export/excel")
def export_excel(db: Session = Depends(get_db), admin: str = Depends(verify_admin)):
    persons = crud.get_authorized_persons(db, limit=1000000)
    wb = Workbook()
    ws = wb.active
    ws.title = "افراد مجاز"
    ws.append(["کدملی", "کدپرسنلی", "نام", "نام خانوادگی", "محل خدمت", "وضعیت خدمت"])
    for p in persons:
        ws.append([
            p.national_code,
            p.personnel_code or "",
            p.first_name,
            p.last_name,
            p.service_location or "",
            p.service_status or ""
        ])
    output = BytesIO()
    wb.save(output)
    output.seek(0)
    return StreamingResponse(
        output,
        media_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
        headers={"Content-Disposition": "attachment; filename=authorized_persons.xlsx"}
    )

# ========== ویرایش/حذف تکی (این روت‌ها باید بعد از روت‌های ثابت بالا تعریف شوند) ==========
@router.put("/{person_id}", response_model=schemas.AuthorizedPerson)
def update_authorized_person(
    person_id: int,
    person: schemas.AuthorizedPersonUpdate,
    db: Session = Depends(get_db),
    admin: str = Depends(verify_admin)
):
    db_obj = crud.update_authorized_person(db, person_id, person)
    if not db_obj:
        raise HTTPException(status_code=404, detail="فرد یافت نشد")
    return db_obj

@router.delete("/{person_id}")
def delete_authorized_person(
    person_id: int,
    db: Session = Depends(get_db),
    admin: str = Depends(verify_admin)
):
    db_obj = crud.delete_authorized_person(db, person_id)
    if not db_obj:
        raise HTTPException(status_code=404, detail="فرد یافت نشد")
    return {"message": "با موفقیت حذف شد"}
